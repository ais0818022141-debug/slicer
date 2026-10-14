import openpyxl

for fname in ["working_on/أسئلة دورات أورام.xlsx", "working_on/أسئلة دورات تجميلية.xlsx"]:
    wb = openpyxl.load_workbook(fname)
    print(f"=== {fname} ===")
    print("Sheets:", wb.sheetnames)
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        print(f"  Sheet {sheet}: max_row={ws.max_row}, max_col={ws.max_column}")
